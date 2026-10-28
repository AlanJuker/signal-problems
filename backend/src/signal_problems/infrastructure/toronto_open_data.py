import csv
import re
from collections.abc import Iterable, Iterator
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

import httpx
from openpyxl import load_workbook

from signal_problems.application.ports import RawDelayRow
from signal_problems.domain.surface import Mode

CKAN_API = "https://ckan0.cf.opendata.inter.prod-toronto.ca/api/3/action"
PACKAGE = "ttc-subway-delay-data"
# The current log is one CSV that grows monthly; 2014–2024 are published as yearly XLSX files.
CURRENT_LOG = "TTC Subway Delay Data since 2025.csv"
HISTORY = re.compile(r"^ttc-subway-delay-(data-)?.*20\d\d$")
# Streetcars and buses: (package, current log). Their 2014–2024 files use another schema (#28).
SURFACE = {
    Mode.STREETCAR: ("ttc-streetcar-delay-data", "TTC Streetcar Delay Data since 2025.csv"),
    Mode.BUS: ("ttc-bus-delay-data", "TTC Bus Delay Data since 2025.csv"),
}
COLUMNS = ("Date", "Time", "Station", "Line", "Code", "Min Delay", "Min Gap", "Bound", "Vehicle")


def parse_csv(lines: Iterable[str]) -> Iterator[RawDelayRow]:
    for row in csv.DictReader(lines):
        yield RawDelayRow(
            date=row["Date"],
            time=row["Time"],
            station=row["Station"],
            line=row["Line"],
            code=row["Code"],
            min_delay=row["Min Delay"],
            min_gap=row["Min Gap"],
            bound=row["Bound"],
            vehicle=row["Vehicle"],
        )


def parse_xlsx(stream: BinaryIO) -> Iterator[RawDelayRow]:
    """Every sheet of a yearly workbook (some years have one sheet per month).

    Cells come typed (dates as datetimes, minutes as numbers); they're turned into the same text
    the CSV has, so both formats go through the same parsing and fingerprinting.
    """
    workbook = load_workbook(stream, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = [str(cell).strip() if cell is not None else "" for cell in next(rows, ())]
            if not set(COLUMNS) <= set(header):
                continue
            index = {name: header.index(name) for name in COLUMNS}
            for cells in rows:
                if not cells or cells[index["Date"]] is None:
                    continue
                value = {name: _text(cells[i]) for name, i in index.items()}
                yield RawDelayRow(
                    date=value["Date"],
                    time=value["Time"],
                    station=value["Station"],
                    line=value["Line"],
                    code=value["Code"],
                    min_delay=value["Min Delay"],
                    min_gap=value["Min Gap"],
                    bound=value["Bound"],
                    vehicle=value["Vehicle"],
                )
    finally:
        workbook.close()


def _text(cell: object) -> str:
    if cell is None:
        return ""
    if isinstance(cell, datetime | date):
        return cell.strftime("%Y-%m-%d")
    if isinstance(cell, time):
        return cell.strftime("%H:%M")
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


class TorontoOpenDataSource:
    """Downloads the delay log through the portal's CKAN API."""

    def __init__(
        self,
        client: httpx.Client,
        *,
        include_history: bool = False,
        package: str = PACKAGE,
        current_log: str = CURRENT_LOG,
    ) -> None:
        self._client = client
        self._include_history = include_history
        self._package = package
        self._current_log = current_log

    def fetch(self) -> Iterator[RawDelayRow]:
        package = self._client.get(f"{CKAN_API}/package_show", params={"id": self._package})
        package.raise_for_status()
        resources = package.json()["result"]["resources"]
        if self._include_history:
            for resource in resources:
                if HISTORY.match(resource["name"]):
                    yield from parse_xlsx(BytesIO(self._download(resource["url"])))
        url = next(r["url"] for r in resources if r["name"] == self._current_log)
        yield from parse_csv(self._download(url).decode("utf-8-sig").splitlines())

    def _download(self, url: str) -> bytes:
        response = self._client.get(url, follow_redirects=True)
        response.raise_for_status()
        return response.content


class CsvFileSource:
    """A delay log saved locally, for offline runs and tests."""

    def __init__(self, path: Path) -> None:
        self._path = path

    def fetch(self) -> Iterator[RawDelayRow]:
        with self._path.open(encoding="utf-8-sig", newline="") as f:
            yield from parse_csv(f)
